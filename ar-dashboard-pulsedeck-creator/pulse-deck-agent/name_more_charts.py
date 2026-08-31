"""One-time: name the remaining large unnamed picture shapes (stale cash/aging
images) on slides 3, 9, 10 so the builder can replace them with placeholders.
Only touches pictures wider than 1.3in that aren't already named chart_/tbl_."""
import os
from pptx import Presentation
from pptx.util import Emu

T = os.path.join(os.path.dirname(os.path.abspath(__file__)), "pulse_template.pptx")
prs = Presentation(T)
slides = list(prs.slides)

# slide index (0-based) -> (prefix, caption, source)
TARGETS = {
    2: ("chart_cash_s3",  "Cash Chart",        "AR Tracker"),   # slide 3
    8: ("chart_cash_s9",  "Cash Scoreboard",   "AR Tracker"),   # slide 9
    9: ("chart_aging_s10","Payments by Timing","YayPay"),       # slide 10
}
named = {}
for idx, (prefix, caption, source) in TARGETS.items():
    i = 0
    for sh in slides[idx].shapes:
        if sh.shape_type != 13:                     # pictures only
            continue
        if sh.name.startswith(("chart_", "tbl_")):  # already handled
            continue
        if Emu(sh.width).inches <= 1.3:             # skip logos / small decor
            continue
        i += 1
        new = f"{prefix}_{i}"
        print(f"  slide {idx+1}: {sh.name!r} ({Emu(sh.width).inches:.1f}x{Emu(sh.height).inches:.1f}in) -> {new}")
        sh.name = new
        named[new] = f"{caption}|{source}"

prs.save(T)
print("\nSaved. Add these to PLACEHOLDER_CHARTS:")
for k, v in named.items():
    print(f'  {k}: "{v}",')

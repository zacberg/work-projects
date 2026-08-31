"""
fix_template.py  — one-time finishing pass on pulse_template.pptx
=================================================================
Builds on the work already done in PowerPoint (chart shapes named, 12 tokens in).
This pass:
  1. gives the 5 data tables stable names the builder targets
  2. tokenizes the 3 literals make_template couldn't (they span runs):
       DSO "37" -> {{DSO}},  "Medium" -> {{FORECAST_CONFIDENCE}},
       write-offs "$0" -> {{AGING90_WRITEOFFS}}

Scoped to specific (slide, table, row, col) cells so nothing else is touched.
Run once:  python fix_template.py
"""
import os
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(HERE, "pulse_template.pptx")

prs = Presentation(TPL)
slides = list(prs.slides)

def tables_on(slide):
    return [sh for sh in slide.shapes if sh.has_table]

# ── 1. name the tables ────────────────────────────────────────────────
# (each slide has exactly one table; match by header text to be safe)
RENAME = {
    1: ("tbl_metrics",  "Result"),       # slide 2 (0-indexed 1)
    3: ("tbl_aging60",  "Aging Controls"),# slide 4
    5: ("tbl_top10",    "Customer"),      # slide 6
    9: ("tbl_aging90",  "Aging Controls"),# slide 10
    10:("tbl_holds",    "Holds"),         # slide 11
}
for idx, (newname, header_hint) in RENAME.items():
    for sh in tables_on(slides[idx]):
        hdr = " ".join(c.text for c in sh.table.rows[0].cells)
        if header_hint in hdr:
            old = sh.name
            sh.name = newname
            print(f"  slide {idx+1}: table {old!r} -> {newname!r}")
            break
    else:
        print(f"  [warn] slide {idx+1}: no table matching header {header_hint!r}")

# ── 2. tokenize the 3 remaining literals (cross-run safe, scoped to cell) ─
def set_cell_token(cell, literal, token):
    """Replace `literal` with `token` inside one table cell, joining runs so
    a value split across runs still gets replaced; keeps run[0] formatting."""
    for p in cell.text_frame.paragraphs:
        full = "".join(r.text for r in p.runs)
        if literal in full and "{{" not in full:
            new = full.replace(literal, token)
            if p.runs:
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
            return True
    return False

def cell_of(slide, table_name, row, col):
    for sh in slide.shapes:
        if sh.has_table and sh.name == table_name:
            return sh.table.cell(row, col)
    return None

fixes = [
    (slides[1], "tbl_metrics", 3, 1, "37", "{{DSO}}"),
    (slides[1], "tbl_metrics", 8, 1, "Medium", "{{FORECAST_CONFIDENCE}}"),
    (slides[9], "tbl_aging90", 4, 1, "$0", "{{AGING90_WRITEOFFS}}"),
]
for slide, tname, r, c, lit, tok in fixes:
    cell = cell_of(slide, tname, r, c)
    if cell is None:
        print(f"  [warn] cell {tname}[{r},{c}] not found")
        continue
    ok = set_cell_token(cell, lit, tok)
    print(f"  {tname}[{r},{c}]: {lit!r} -> {tok}   {'OK' if ok else 'NOT FOUND (already token?)'}")

prs.save(TPL)
print("\nSaved pulse_template.pptx")

# verify
chk = Presentation(TPL)
names = [sh.name for s in chk.slides for sh in s.shapes if sh.has_table]
print("Table names now:", names)

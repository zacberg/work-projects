"""
name_charts.py — give the chart pictures stable names so build_pulse.py / the agent
can target each for image-swap. Decorative images (stock photos, icons) left alone.

Usage:
  python name_charts.py          # dry run: report what WOULD be named
  python name_charts.py save     # actually write names (close PowerPoint first!)
"""
import os, sys, hashlib, zipfile
from pptx import Presentation

T = os.path.join(os.path.dirname(os.path.abspath(__file__)), "pulse_template.pptx")
SAVE = len(sys.argv) > 1 and sys.argv[1] == "save"

MEDIA_TO_NAME = {
    "image15.png": "chart_weekly_run_rate",
    "image18.png": "chart_total_ar_aging",
    "image19.png": "chart_gt60_trend",
    "image20.png": "chart_upcoming_to_60",
    "image22.png": "chart_root_cause",
    "image24.png": "chart_payment_methods",
    "image27.png": "chart_payments_against_60",
    "image29.png": "chart_gt90_trend",
    "image30.png": "chart_payments_by_timing_pct",
    "image31.png": "chart_payments_by_timing_dollars",
    "image13.png": "chart_cash_target_vs_actual",
}

# sha1 -> media filename (so we can identify each picture reliably)
z = zipfile.ZipFile(T)
sha_to_file = {}
for n in z.namelist():
    if n.startswith("ppt/media/"):
        sha_to_file[hashlib.sha1(z.read(n)).hexdigest()] = n.split("/")[-1]

prs = Presentation(T)
used = {}
named = 0
print("PICTURE INVENTORY (slide : media : name)\n")
for i, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        if shape.__class__.__name__ != "Picture":
            continue
        try:
            mf = sha_to_file.get(shape.image.sha1)
        except Exception:
            mf = None
        if mf in MEDIA_TO_NAME:
            base = MEDIA_TO_NAME[mf]
            used[base] = used.get(base, 0) + 1
            name = base if used[base] == 1 else f"{base}_{used[base]}"
            shape.name = name
            named += 1
            print(f"  slide{i:<2} {mf:<14} -> {name}   *CHART*")
        else:
            print(f"  slide{i:<2} {mf or '(emf/other)':<14}    (decorative — left as-is)")

print(f"\n{named} chart pictures identified.")
if SAVE:
    prs.save(T)
    print("SAVED names into pulse_template.pptx")
else:
    print("DRY RUN — no changes written. Re-run with 'save' (PowerPoint must be closed).")

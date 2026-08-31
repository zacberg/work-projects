import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import datetime, random, os

random.seed(99)
output_dir = r"C:\Users\ZachBergman\Desktop\BackPocketBlue_TestInputs"
os.makedirs(output_dir, exist_ok=True)

# ── Product → LOB + GL + PH05 + Category mappings (from real file) ─────────────
products = [
    # (PH06, PH05, LOB,               GL,    Category, SW_SRVC)
    ("INF",  "DDI", "Distribution",    40100, "Maintain",  "Software"),
    ("ERO",  "D1S", "Distribution",    40100, "Maintain",  "Software"),
    ("COS",  "COM", "Distribution",    40100, "Invest",    "Software"),
    ("PEP",  "PEP", "Distribution",    40100, "Invest",    "Software"),
    ("CVE",  "CVO", "Distribution",    40100, "Invest",    "Software"),
    ("OPM",  "CVO", "Distribution",    40100, "Invest",    "Software"),
    ("LUC",  "CVO", "Distribution",    40100, "Invest",    "Software"),
    ("KWP",  "KWP", "Packaging",       40200, "Maintain",  "Software"),
    ("P3K",  "ABS", "Packaging",       40200, "Invest",    "Software"),
    ("ENA",  "IQS", "Quality",         40400, "Invest",    "Software"),
    ("POD",  "IQS", "Quality",         40400, "Invest",    "Software"),
    ("POP",  "IQS", "Quality",         40400, "Invest",    "Software"),
    ("GAG",  "PQS", "Quality",         40400, "Maintain",  "Software"),
    ("SQC",  "PQS", "Quality",         40400, "Maintain",  "Software"),
    ("PRO",  "PRO", "Manufacturing",   40500, "Maintain",  "Software"),
    ("SMS",  "PIN", "Manufacturing",   40500, "Invest",    "Software"),
    ("PMG",  "PFC", "Customer Focused",40600, "CFP",       "Software"),
    ("SPC",  "DQS", "Customer Focused",40600, "CFP",       "Software"),
    ("MAN",  "VIA", "Customer Focused",40600, "CFP",       "Software"),
    ("VER",  "VER", "Customer Focused",40600, "CFP",       "Software"),
    ("ADW",  "ADW", "Customer Focused",40600, "CFP",       "Software"),
]

sources = ["Current", "Budget", "June BOD Forecast"]
measures = ["ACV Bookings", "ARR Bookings", "LEB SW Bookings"]
rev_type_by_measure = {
    "ACV Bookings":     ["Services", "Subscription", "Term License"],
    "ARR Bookings":     ["Subscription"],
    "LEB SW Bookings":  ["Subscription", "Term License"],
}
rec_by_rev = {
    "Services":        "Non-Recurring Revenue",
    "Subscription":    "Recurring Revenue",
    "Term License":    "Recurring Revenue",
    "License":         "Non-Recurring Revenue",
    "Maintenance":     "Recurring Revenue",
    "Other Recurring": "Recurring Revenue",
    "Hardware":        "Non-Recurring Revenue",
}

rows = []
for source in sources:
    for year in [2025, 2026]:
        for quarter in [1, 2, 3, 4]:
            months = {1: [1,2,3], 2: [4,5,6], 3: [7,8,9], 4: [10,11,12]}[quarter]
            h1h2  = "1H" if quarter <= 2 else "2H"
            for month in months:
                ytd_roy = "YTD" if (year == 2026 and quarter <= 2) else "ROY"
                date = datetime.datetime(year, month, 1)
                for ph6, ph5, lob, gl, cat, sw in products:
                    for measure in measures:
                        for rev_type in rev_type_by_measure[measure]:
                            base = random.uniform(0.5, 25.0)
                            if rev_type == "Services":
                                base *= 0.4
                            compare_val = round(base * random.uniform(0.85, 1.15), 6)
                            rows.append({
                                "Source":           source,
                                "GL":               gl,
                                "Product":          ph6,
                                "Revenue Type":     rev_type,
                                "Date":             date,
                                "Measure":          measure,
                                "Value":            round(base, 6),
                                "Year":             year,
                                "Quarter":          quarter,
                                "Month":            month,
                                "1H_2H":            h1h2,
                                "YTD_ROY":          ytd_roy,
                                "Compare Value":    compare_val,
                                "LOB":              lob,
                                "Product_Category": cat,
                                "PH05":             ph5,
                                "SW_SRVC":          sw,
                                "Rec_Nrec":         rec_by_rev[rev_type],
                            })

df_book = pd.DataFrame(rows)

wb_b = Workbook()
ws_b = wb_b.active
ws_b.title = "Bookings.Detail"
headers = list(df_book.columns)
hfill  = PatternFill("solid", fgColor="3B2845")
hfont  = Font(bold=True, color="FFFFFF", name="Calibri")
for c, h in enumerate(headers, 1):
    cell = ws_b.cell(row=1, column=c, value=h)
    cell.fill, cell.font = hfill, hfont
for r, row in df_book.iterrows():
    for c, h in enumerate(headers, 1):
        ws_b.cell(row=r+2, column=c, value=row[h])
for col in ws_b.columns:
    ws_b.column_dimensions[col[0].column_letter].width = max(
        len(str(c.value)) if c.value else 0 for c in col) + 3

book_path = os.path.join(output_dir, "Bookings_Input_072226.xlsx")
wb_b.save(book_path)
print(f"Saved: {book_path}  ({len(df_book)} rows)")

# ── Capacity_Input ──────────────────────────────────────────────────────────────
ph_map = [
    # (PH04,  PH05, PH06)
    ("DIST", "DDI", "INF"),
    ("DIST", "D1S", "ERO"),
    ("DIST", "COM", "COS"),
    ("DIST", "PEP", "PEP"),
    ("DIST", "CVO", "CVE"),
    ("PACK", "KWP", "KWP"),
    ("PACK", "ABS", "P3K"),
    ("QUAL", "IQS", "ENA"),
    ("QUAL", "IQS", "POD"),
    ("QUAL", "IQS", "POP"),
    ("QUAL", "PQS", "GAG"),
    ("QUAL", "PQS", "SQC"),
    ("MANF", "PRO", "PRO"),
    ("MANF", "PIN", "SMS"),
    ("CFP",  "ADW", "ADW"),
    ("CFP",  "DQS", "SPC"),
    ("CFP",  "VIA", "MAN"),
    ("CFP",  "VER", "VER"),
    ("CFP",  "PFC", "PMG"),
]

quarter_months = {
    "Q1-26": [(2026,1,31),(2026,2,28),(2026,3,31)],
    "Q2-26": [(2026,4,30),(2026,5,31),(2026,6,30)],
    "Q3-26": [(2026,7,31),(2026,8,31),(2026,9,30)],
    "Q4-26": [(2026,10,31),(2026,11,30),(2026,12,31)],
}

left_rows, right_rows = [], []
ph4_hc = {}  # track beginning HC per PH04+quarter
for qtr, months in quarter_months.items():
    for ph4, ph5, ph6 in ph_map:
        key = (ph4, qtr)
        if key not in ph4_hc:
            ph4_hc[key] = random.randint(10, 30)
        for y, m, d in months:
            date = datetime.datetime(y, m, d)
            eff_hc     = round(ph4_hc[key] * random.uniform(0.9, 1.1), 6)
            quota      = round(random.uniform(200, 800), 6)
            exp_book   = round(quota * random.uniform(0.7, 1.1), 6)
            left_rows.append((ph4, ph5, ph6, qtr, date, eff_hc, quota, exp_book))

            beg_hc    = ph4_hc[key]
            hires     = random.randint(0, 3)
            exits     = random.randint(0, 2)
            transfers = random.randint(-1, 1)
            right_rows.append((ph4, qtr, date, beg_hc, hires, exits, transfers))

wb_c = Workbook()
ws_c = wb_c.active
ws_c.title = "Calculations"

# Row 1 & 2: title rows matching real file
ws_c.cell(1,  1, "Bookings/Quota by Product & Month")
ws_c.cell(1, 12, "Headcount by BU & Month")
ws_c.cell(2,  1, "Source: CapModel tab from Capacity_Model_Current")
ws_c.cell(2, 12, "Source: CapModel tab from Capacity_Model_Current")

# Row 3: headers
left_h  = ["PH04","PH05","PH06","Quarter","Month","Effective HC","Quota","ExpBookings"]
right_h = ["PH04.1","Quarter.1","Month.1","Beginning HC","Hires","Exits","Transfers"]
for c, h in enumerate(left_h, 1):
    cell = ws_c.cell(3, c, h)
    cell.fill, cell.font = hfill, hfont
for c, h in enumerate(right_h, 12):
    cell = ws_c.cell(3, c, h)
    cell.fill, cell.font = hfill, hfont

# Data rows starting at row 4
for i, (lr, rr) in enumerate(zip(left_rows, right_rows)):
    r = i + 4
    for c, v in enumerate(lr, 1):
        ws_c.cell(r, c, v)
    for c, v in enumerate(rr, 12):
        ws_c.cell(r, c, v)

for col in ws_c.columns:
    ws_c.column_dimensions[col[0].column_letter].width = 14

cap_path = os.path.join(output_dir, "Capacity_Input_072226.xlsx")
wb_c.save(cap_path)
print(f"Saved: {cap_path}  ({len(left_rows)} rows)")
print(f"\nFolder: {output_dir}")
